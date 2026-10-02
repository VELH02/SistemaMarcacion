import os
import time
import multiprocessing as mp
import tkinter as tk
from tkinter import ttk
from datetime import datetime
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import numpy as np
import cv2
import face_recognition
from PIL import Image, ImageTk

# ----------------------------------------------------------------------
# CONFIGURACIÓN DE ARCHIVOS Y PALETA DE COLORES (VERDE NUTRESA)
# ----------------------------------------------------------------------
EXCEL_MARCACIONES = "marcaciones.xlsx"
EXCEL_EMPLEADOS = "empleados.xlsx"
DIR_ROSTROS = "base_rostros"
HOJA_REGISTRO = "Registro"
ADMIN_PIN = "123456"

# Umbral de distancia euclidiana: Menor valor = Mayor rigurosidad
UMBRAL_DISTANCIA_ESTRICTO = 0.48  

FONT_PRIMARY = "Segoe UI"

COLOR_PRIMARY = "#008037"        # Verde Oficial Nutresa
COLOR_PRIMARY_HOVER = "#00642B"  # Verde Oscuro (Hover)
COLOR_SECONDARY = "#2D3732"      # Gris Bosque Oscuro
COLOR_SECONDARY_HOVER = "#1E2622"
COLOR_BG = "#F2F5F3"             # Fondo General
COLOR_CARD = "#FFFFFF"           # Tarjetas / Contenedores
COLOR_TEXT_DARK = "#1A241E"      # Texto Principal
COLOR_MUTED = "#637368"          # Texto Secundario
COLOR_BORDER = "#D3DED7"         # Bordes
COLOR_DANGER = "#D9383A"         # Rojo Alerta / Borrar
COLOR_WARNING = "#D9822B"        # Naranja Advertencia


# ----------------------------------------------------------------------
# DETECCIÓN DE PARPADEO (EYE ASPECT RATIO - EAR)
# ----------------------------------------------------------------------
def calcular_ear(ojo):
    """Calcula la relación de aspecto del ojo para detectar parpadeos."""
    A = np.linalg.norm(ojo[1] - ojo[5])
    B = np.linalg.norm(ojo[2] - ojo[4])
    C = np.linalg.norm(ojo[0] - ojo[3])
    ear = (A + B) / (2.0 * C)
    return ear


# ----------------------------------------------------------------------
# TRABAJADOR EN PROCESO INDEPENDIENTE (IA + CÁLCULO DE LIVENESS)
# ----------------------------------------------------------------------
def worker_reconocimiento(input_queue, output_queue):
    cont_parpadeos = 0
    ojo_cerrado = False

    while True:
        try:
            item = input_queue.get()
            if item is None:
                break

            frame_rgb, vector_registrado = item
            frame_h, frame_w, _ = frame_rgb.shape

            frame_small = cv2.resize(frame_rgb, (0, 0), fx=0.5, fy=0.5)
            locations_small = face_recognition.face_locations(frame_small, model="hog")

            if len(locations_small) == 0:
                output_queue.put({
                    "rect": None, "verificado": False, "similitud": 0.0,
                    "alerta_liveness": "Buscando rostro..."
                })
                continue

            top, right, bottom, left = locations_small[0]
            l, t, r, b = left * 2, top * 2, right * 2, bottom * 2
            l, t, r, b = max(0, l), max(0, t), min(frame_w, r), min(frame_h, b)

            landmarks_list = face_recognition.face_landmarks(frame_rgb, [(t, r, b, l)])
            
            if len(landmarks_list) > 0:
                landmarks = landmarks_list[0]
                ojo_izq = np.array(landmarks["left_eye"])
                ojo_der = np.array(landmarks["right_eye"])

                ear_izq = calcular_ear(ojo_izq)
                ear_der = calcular_ear(ojo_der)
                ear_promedio = (ear_izq + ear_der) / 2.0

                if ear_promedio < 0.20:
                    ojo_cerrado = True
                elif ear_promedio > 0.25 and ojo_cerrado:
                    cont_parpadeos += 1
                    ojo_cerrado = False

            if cont_parpadeos < 1:
                output_queue.put({
                    "rect": (l, t, r, b), "verificado": False, "similitud": 0.0,
                    "alerta_liveness": "👁 PARPADEE PARA PROBAR QUE ES REAL"
                })
                continue

            encodings = face_recognition.face_encodings(frame_rgb, [(t, r, b, l)])
            if len(encodings) > 0:
                encoding_capturado = encodings[0]
                distancia = face_recognition.face_distance([vector_registrado], encoding_capturado)[0]
                es_valido = distancia < UMBRAL_DISTANCIA_ESTRICTO
                similitud = max(0.0, min(100.0, (1.0 - (distancia / UMBRAL_DISTANCIA_ESTRICTO)) * 100))

                output_queue.put({
                    "rect": (l, t, r, b),
                    "verificado": es_valido,
                    "similitud": similitud,
                    "alerta_liveness": ""
                })
            else:
                output_queue.put({
                    "rect": (l, t, r, b), "verificado": False, "similitud": 0.0,
                    "alerta_liveness": ""
                })

        except Exception:
            pass


# ----------------------------------------------------------------------
# FORMATO Y ESTILOS PARA LA HOJA CONSOLIDADA DE EXCEL
# ----------------------------------------------------------------------
def aplicar_formato_excel(ws):
    """Aplica formato corporativo estilizado a la hoja de registro consolidada."""
    fill_header = PatternFill(start_color="008037", end_color="008037", fill_type="solid")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    border_thin = Border(
        left=Side(style='thin', color='D3DED7'),
        right=Side(style='thin', color='D3DED7'),
        top=Side(style='thin', color='D3DED7'),
        bottom=Side(style='thin', color='D3DED7')
    )
    
    # Encabezados
    ws.row_dimensions[1].height = 26
    for col_num in range(1, 10):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    # Anchos de columna óptimos
    anchos = {
        "A": 16, # ID Colaborador
        "B": 20, # Departamento
        "C": 14, # Fecha
        "D": 15, # Hora Entrada
        "E": 15, # Hora Salida
        "F": 18, # Horas Trabajadas
        "G": 16, # Minutos Tarde
        "H": 16, # Tiempo Extra
        "I": 28  # Notas
    }
    for col_letter, width in anchos.items():
        ws.column_dimensions[col_letter].width = width

    # Filas de datos
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 20
        for c in range(1, 10):
            cell = ws.cell(row=r, column=c)
            cell.border = border_thin
            cell.font = Font(name="Segoe UI", size=10)
            if c in [1, 3, 4, 5, 6, 7, 8]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left


# ----------------------------------------------------------------------
# MANEJO DE ARCHIVOS DE EXCEL Y CARPETAS
# ----------------------------------------------------------------------
def inicializar_entorno():
    if not os.path.exists(DIR_ROSTROS):
        os.makedirs(DIR_ROSTROS)

    headers = ["ID Colaborador", "Departamento", "Fecha", "Hora Entrada", "Hora Salida", "Horas Trabajadas", "Minutos Tarde", "Tiempo Extra", "Notas"]

    if not os.path.exists(EXCEL_MARCACIONES):
        wb = Workbook()
        ws = wb.active
        ws.title = HOJA_REGISTRO
        ws.append(headers)
        aplicar_formato_excel(ws)
        wb.save(EXCEL_MARCACIONES)
    else:
        wb = load_workbook(EXCEL_MARCACIONES)
        if HOJA_REGISTRO not in wb.sheetnames:
            ws = wb.create_sheet(title=HOJA_REGISTRO)
            ws.append(headers)
            aplicar_formato_excel(ws)
            wb.save(EXCEL_MARCACIONES)

    if not os.path.exists(EXCEL_EMPLEADOS):
        wb = Workbook()
        ws = wb.active
        ws.title = "Empleados"
        ws.append(["ID_Empleado", "Departamento", "Archivo_Vector", "Estado"])
        wb.save(EXCEL_EMPLEADOS)


def existe_empleado(emp_id):
    if not os.path.exists(EXCEL_EMPLEADOS):
        return False, ""
    wb = load_workbook(EXCEL_EMPLEADOS)
    if "Empleados" not in wb.sheetnames:
        return False, ""
    ws = wb["Empleados"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None and str(row[0]).strip() == str(emp_id).strip():
            depto = str(row[1]) if len(row) > 1 and row[1] is not None else "N/A"
            return True, depto
    return False, ""


def guardar_nuevo_empleado(emp_id, departamento, archivo_vector):
    wb_emp = load_workbook(EXCEL_EMPLEADOS)
    if "Empleados" not in wb_emp.sheetnames:
        ws_emp = wb_emp.create_sheet("Empleados")
        ws_emp.append(["ID_Empleado", "Departamento", "Archivo_Vector", "Estado"])
    else:
        ws_emp = wb_emp["Empleados"]
    ws_emp.append([str(emp_id), departamento, archivo_vector, "ACTIVO"])
    wb_emp.save(EXCEL_EMPLEADOS)

    inicializar_entorno()


def registrar_marcacion_excel(emp_id, tipo_marcacion, nota=""):
    """
    Registra marcaciones de todos los colaboradores en la única hoja 'Registro'.
    Preserva las demás hojas (incluyendo 'Base Legal').
    Garantiza 1 sola fila por colaborador y por fecha:
    - Entrada: Mantiene la hora más temprana.
    - Salida: Mantiene la hora más tardía.
    Inyecta fórmulas de cálculo automático de horas.
    """
    inicializar_entorno()

    wb = load_workbook(EXCEL_MARCACIONES)
    
    if HOJA_REGISTRO not in wb.sheetnames:
        ws = wb.create_sheet(title=HOJA_REGISTRO)
        ws.append(["ID Colaborador", "Departamento", "Fecha", "Hora Entrada", "Hora Salida", "Horas Trabajadas", "Minutos Tarde", "Tiempo Extra", "Notas"])
    else:
        ws = wb[HOJA_REGISTRO]

    _, depto = existe_empleado(emp_id)
    if not depto:
        depto = "N/A"

    ahora = datetime.now()
    fecha_hoy = ahora.strftime("%Y-%m-%d")
    hora_actual = ahora.strftime("%H:%M:%S")

    # Buscar si ya existe registro para este colaborador en el día de hoy
    fila_encontrada = None
    for r in range(2, ws.max_row + 1):
        val_id = ws.cell(row=r, column=1).value
        val_fecha = ws.cell(row=r, column=3).value
        if val_id is not None and str(val_id).strip() == str(emp_id).strip() and val_fecha is not None and str(val_fecha).strip() == fecha_hoy:
            fila_encontrada = r
            break

    if fila_encontrada is not None:
        r = fila_encontrada
        if tipo_marcacion == "ENTRADA":
            val_actual = ws.cell(row=r, column=4).value
            val_str = str(val_actual).strip() if val_actual is not None else ""
            if not val_str or val_str == "None" or hora_actual < val_str:
                ws.cell(row=r, column=4, value=hora_actual)

        elif tipo_marcacion in ["SALIDA", "TIEMPO EXTRA"]:
            val_actual = ws.cell(row=r, column=5).value
            val_str = str(val_actual).strip() if val_actual is not None else ""
            if not val_str or val_str == "None" or hora_actual > val_str:
                ws.cell(row=r, column=5, value=hora_actual)

        if nota:
            nota_previa = ws.cell(row=r, column=9).value
            nueva_nota = f"{nota_previa}; {nota}" if nota_previa else nota
            ws.cell(row=r, column=9, value=nueva_nota)

    else:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=str(emp_id))
        ws.cell(row=r, column=2, value=depto)
        ws.cell(row=r, column=3, value=fecha_hoy)

        if tipo_marcacion == "ENTRADA":
            ws.cell(row=r, column=4, value=hora_actual)
            ws.cell(row=r, column=5, value="")
        else:
            ws.cell(row=r, column=4, value="")
            ws.cell(row=r, column=5, value=hora_actual)

        ws.cell(row=r, column=9, value=nota if nota else "")

    # Aplicar fórmulas dinámicas de Excel para cálculos de horas
    ws.cell(row=r, column=6, value=f'=IF(AND(D{r}<>"",E{r}<>""),(E{r}-D{r})*24,"")')
    ws.cell(row=r, column=7, value=f'=IF(AND(D{r}<>"",D{r}>TIME(8,0,0)),(D{r}-TIME(8,0,0))*1440,0)')
    ws.cell(row=r, column=8, value=f'=IF(AND(F{r}<>"",F{r}>8),F{r}-8,0)')

    aplicar_formato_excel(ws)
    wb.save(EXCEL_MARCACIONES)


def obtener_historial_empleado(emp_id):
    if not os.path.exists(EXCEL_MARCACIONES):
        return []
    wb = load_workbook(EXCEL_MARCACIONES, data_only=True)
    if HOJA_REGISTRO not in wb.sheetnames:
        return []

    ws = wb[HOJA_REGISTRO]
    registros = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None and str(row[0]).strip() == str(emp_id).strip():
            fecha = str(row[2]) if len(row) > 2 and row[2] is not None else "--"
            entrada = str(row[3]) if len(row) > 3 and row[3] is not None else "--:--:--"
            salida = str(row[4]) if len(row) > 4 and row[4] is not None else "--:--:--"
            horas = str(row[5]) if len(row) > 5 and row[5] is not None else "-"
            tarde = str(row[6]) if len(row) > 6 and row[6] is not None else "-"
            
            registros.append({
                "fecha": fecha,
                "entrada": entrada,
                "salida": salida,
                "horas": horas,
                "tarde": tarde
            })
    return registros


# ----------------------------------------------------------------------
# COMPONENTE ENCABEZADO
# ----------------------------------------------------------------------
class HeaderNutresa(tk.Frame):
    def __init__(self, parent):
        super().__init__(parent, bg=COLOR_BG)
        
        top_accent = tk.Frame(self, bg=COLOR_PRIMARY, height=6)
        top_accent.pack(fill="x", side="top")

        brand_frame = tk.Frame(self, bg=COLOR_BG)
        brand_frame.pack(pady=(10, 5))

        if os.path.exists("logo.png"):
            try:
                img_pil = Image.open("logo.png")
                img_pil.thumbnail((220, 50), Image.Resampling.LANCZOS)
                self.img = ImageTk.PhotoImage(img_pil)
                lbl_logo = tk.Label(brand_frame, image=self.img, bg=COLOR_BG)
                lbl_logo.pack()
            except Exception:
                self._crear_logo_texto(brand_frame)
        else:
            self._crear_logo_texto(brand_frame)

    def _crear_logo_texto(self, parent):
        lbl_brand = tk.Label(parent, text="Grupo Nutresa", font=(FONT_PRIMARY, 20, "bold"), fg=COLOR_PRIMARY, bg=COLOR_BG)
        lbl_brand.pack()
        lbl_sub = tk.Label(parent, text="SISTEMA DE ASISTENCIA Y TIEMPO", font=(FONT_PRIMARY, 8, "bold"), fg=COLOR_MUTED, bg=COLOR_BG)
        lbl_sub.pack(pady=(2, 0))


# ----------------------------------------------------------------------
# APLICACIÓN PRINCIPAL
# ----------------------------------------------------------------------
class SistemaMarcacionApp(tk.Tk):
    def __init__(self):
        super().__init__()
        
        self.title("Grupo Nutresa - Control de Asistencia Biométrico")
        self.geometry("840x670")
        self.minsize(800, 600)
        self.configure(bg=COLOR_BG)

        inicializar_entorno()

        self.current_emp_code = ""
        self.current_action = ""
        self.timeout_seconds = 30
        self.timer_id = None

        self.container = tk.Frame(self, bg=COLOR_BG)
        self.container.pack(fill="both", expand=True)

        self.bind_all("<Button-1>", self.reset_inactivity_timer)
        self.bind_all("<Key>", self.reset_inactivity_timer)

        self.frames = {}
        for F in (PantallaPrincipal, PantallaTeclado, PantallaBiometrica, PantallaHistorial, PantallaEnrolar, PantallaAdminPin):
            page_name = F.__name__
            frame = F(parent=self.container, controller=self)
            self.frames[page_name] = frame
            frame.grid(row=0, column=0, sticky="nsew")

        self.container.grid_rowconfigure(0, weight=1)
        self.container.grid_columnconfigure(0, weight=1)

        self.show_frame("PantallaPrincipal")

    def show_frame(self, page_name):
        for name, frame in self.frames.items():
            if hasattr(frame, "detener_camara"):
                frame.detener_camara()

        frame = self.frames[page_name]
        frame.tkraise()
        if hasattr(frame, "on_show"):
            frame.on_show()
        self.reset_inactivity_timer()

    def reset_inactivity_timer(self, event=None):
        if self.timer_id:
            self.after_cancel(self.timer_id)
        
        if self.get_current_frame_name() != "PantallaPrincipal":
            self.timer_id = self.after(self.timeout_seconds * 1000, self.auto_return_to_main)

    def auto_return_to_main(self):
        self.current_emp_code = ""
        self.current_action = ""
        self.show_frame("PantallaPrincipal")

    def get_current_frame_name(self):
        for name, frame in self.frames.items():
            if frame.winfo_viewable():
                return name
        return "PantallaPrincipal"


# ----------------------------------------------------------------------
# VISTA: PANTALLA PRINCIPAL
# ----------------------------------------------------------------------
class PantallaPrincipal(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg=COLOR_BG)
        self.controller = controller

        HeaderNutresa(self).pack(fill="x")

        lbl_instruccion = tk.Label(
            self, text="Seleccione la operación que desea registrar", 
            font=(FONT_PRIMARY, 12, "bold"), fg=COLOR_MUTED, bg=COLOR_BG
        )
        lbl_instruccion.pack(pady=(10, 5))

        center_card = tk.Frame(self, bg=COLOR_CARD, bd=1, relief="solid", highlightthickness=0)
        center_card.pack(expand=True, padx=70, pady=(5, 15), fill="both")

        grid_frame = tk.Frame(center_card, bg=COLOR_CARD)
        grid_frame.pack(expand=True, fill="both", padx=20, pady=20)
        
        grid_frame.columnconfigure(0, weight=1, pad=12)
        grid_frame.columnconfigure(1, weight=1, pad=12)
        grid_frame.rowconfigure(0, weight=1, pad=12)
        grid_frame.rowconfigure(1, weight=1, pad=12)

        self._crear_boton_estilizado(grid_frame, "📥  ENTRADA", "Registro de Ingreso", COLOR_PRIMARY, COLOR_PRIMARY_HOVER, 0, 0, "ENTRADA")
        self._crear_boton_estilizado(grid_frame, "📤  SALIDA", "Registro de Egreso", COLOR_SECONDARY, COLOR_SECONDARY_HOVER, 0, 1, "SALIDA")
        self._crear_boton_estilizado(grid_frame, "📋  MARCACIONES", "Consultar Historial", COLOR_SECONDARY, COLOR_SECONDARY_HOVER, 1, 0, "MARCACIONES")
        self._crear_boton_estilizado(grid_frame, "⏱  HORA EXTRA", "Sobretiempo Especial", COLOR_SECONDARY, COLOR_SECONDARY_HOVER, 1, 1, "TIEMPO EXTRA")

        btn_enrolar = tk.Button(
            self, text="⚙️  ENROLAR NUEVO COLABORADOR (ADMIN)", font=(FONT_PRIMARY, 9, "bold"),
            bg=COLOR_BG, fg=COLOR_MUTED, activebackground=COLOR_BORDER, bd=0, relief="flat", cursor="hand2",
            command=lambda: self.controller.show_frame("PantallaAdminPin")
        )
        btn_enrolar.pack(pady=(0, 15))

    def _crear_boton_estilizado(self, parent, titulo, subtitulo, bg, bg_hover, row, col, accion):
        btn = tk.Button(
            parent, text=f"{titulo}\n{subtitulo}", font=(FONT_PRIMARY, 12, "bold"),
            bg=bg, fg="#FFFFFF", activebackground=bg_hover, activeforeground="#FFFFFF",
            bd=0, relief="flat", cursor="hand2", justify="center",
            command=lambda: self.seleccionar_accion(accion)
        )
        btn.bind("<Enter>", lambda e: btn.config(bg=bg_hover))
        btn.bind("<Leave>", lambda e: btn.config(bg=bg))
        btn.grid(row=row, column=col, sticky="nsew", ipadx=10, ipady=10)

    def seleccionar_accion(self, accion):
        self.controller.current_action = accion
        self.controller.current_emp_code = ""
        self.controller.show_frame("PantallaTeclado")


# ----------------------------------------------------------------------
# VISTA: ACCESO ADMINISTRATIVO
# ----------------------------------------------------------------------
class PantallaAdminPin(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg=COLOR_BG)
        self.controller = controller

        HeaderNutresa(self).pack(fill="x")

        card = tk.Frame(self, bg=COLOR_CARD, bd=1, relief="solid")
        card.pack(pady=10, padx=140, fill="both", expand=True)

        lbl_title = tk.Label(card, text="ACCESO ADMINISTRATIVO", font=(FONT_PRIMARY, 12, "bold"), fg=COLOR_TEXT_DARK, bg=COLOR_CARD)
        lbl_title.pack(pady=(10, 2))

        self.display_var = tk.StringVar()
        self.display = tk.Entry(
            card, textvariable=self.display_var, show="*", font=(FONT_PRIMARY, 18, "bold"),
            justify="center", bd=2, relief="solid", state="readonly", 
            fg=COLOR_TEXT_DARK, bg=COLOR_CARD
        )
        self.display.pack(pady=5, ipadx=10, ipady=4)

        self.lbl_status = tk.Label(card, text="", font=(FONT_PRIMARY, 9, "bold"), bg=COLOR_CARD)
        self.lbl_status.pack(pady=2)

        keypad = tk.Frame(card, bg=COLOR_CARD)
        keypad.pack(pady=5)

        buttons = [
            ('7', 0, 0), ('8', 0, 1), ('9', 0, 2),
            ('4', 1, 0), ('5', 1, 1), ('6', 1, 2),
            ('1', 2, 0), ('2', 2, 1), ('3', 2, 2),
            ('BORRAR', 3, 0), ('0', 3, 1), ('ENTRAR', 3, 2)
        ]

        for text, row, col in buttons:
            if text == 'BORRAR':
                btn = tk.Button(keypad, text=text, font=(FONT_PRIMARY, 9, "bold"), bg=COLOR_DANGER, fg="#FFFFFF", width=7, height=2, bd=0, relief="flat", cursor="hand2", command=self.press_borrar)
            elif text == 'ENTRAR':
                btn = tk.Button(keypad, text=text, font=(FONT_PRIMARY, 9, "bold"), bg=COLOR_PRIMARY, fg="#FFFFFF", width=7, height=2, bd=0, relief="flat", cursor="hand2", command=self.verificar_pin)
            else:
                btn = tk.Button(keypad, text=text, font=(FONT_PRIMARY, 12, "bold"), bg=COLOR_BG, fg=COLOR_TEXT_DARK, width=5, height=2, bd=1, relief="solid", cursor="hand2", command=lambda t=text: self.press_num(t))
            btn.grid(row=row, column=col, padx=3, pady=3)

        btn_cancel = tk.Button(card, text="CANCELAR", font=(FONT_PRIMARY, 9, "bold"), bg=COLOR_SECONDARY, fg="#FFFFFF", bd=0, relief="flat", cursor="hand2", command=lambda: controller.show_frame("PantallaPrincipal"))
        btn_cancel.pack(pady=8, ipadx=15, ipady=4)

    def on_show(self):
        self.display_var.set("")
        self.lbl_status.config(text="")

    def press_num(self, num):
        self.lbl_status.config(text="")
        curr = self.display_var.get()
        if len(curr) < 8:
            self.display_var.set(curr + num)

    def press_borrar(self):
        self.lbl_status.config(text="")
        self.display_var.set(self.display_var.get()[:-1])

    def verificar_pin(self):
        if self.display_var.get() == ADMIN_PIN:
            self.controller.show_frame("PantallaEnrolar")
        else:
            self.lbl_status.config(text="❌ PIN Incorrecto", fg=COLOR_DANGER)
            self.display_var.set("")


# ----------------------------------------------------------------------
# VISTA: TECLADO NUMÉRICO TÁCTIL
# ----------------------------------------------------------------------
class PantallaTeclado(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg=COLOR_BG)
        self.controller = controller

        HeaderNutresa(self).pack(fill="x")

        self.lbl_accion = tk.Label(self, text="", font=(FONT_PRIMARY, 12, "bold"), fg=COLOR_PRIMARY, bg=COLOR_BG)
        self.lbl_accion.pack(pady=5)

        self.display_var = tk.StringVar()
        self.display = tk.Entry(
            self, textvariable=self.display_var, font=(FONT_PRIMARY, 22, "bold"),
            justify="center", bd=2, relief="solid", state="readonly", 
            fg=COLOR_TEXT_DARK, bg=COLOR_CARD, highlightthickness=1, highlightcolor=COLOR_PRIMARY
        )
        self.display.pack(pady=5, ipadx=10, ipady=6)

        self.lbl_msg = tk.Label(self, text="", font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_BG)
        self.lbl_msg.pack(pady=2)

        card_keypad = tk.Frame(self, bg=COLOR_CARD, bd=1, relief="solid")
        card_keypad.pack(pady=5, padx=120)

        keypad = tk.Frame(card_keypad, bg=COLOR_CARD)
        keypad.pack(padx=15, pady=10)

        buttons = [
            ('7', 0, 0), ('8', 0, 1), ('9', 0, 2),
            ('4', 1, 0), ('5', 1, 1), ('6', 1, 2),
            ('1', 2, 0), ('2', 2, 1), ('3', 2, 2),
            ('BORRAR', 3, 0), ('0', 3, 1), ('ACEPTAR', 3, 2)
        ]

        for text, row, col in buttons:
            if text == 'BORRAR':
                btn = tk.Button(keypad, text=text, font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_DANGER, fg="#FFFFFF", width=8, height=2, bd=0, relief="flat", cursor="hand2", command=self.press_borrar)
            elif text == 'ACEPTAR':
                btn = tk.Button(keypad, text=text, font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_PRIMARY, fg="#FFFFFF", width=8, height=2, bd=0, relief="flat", cursor="hand2", command=self.press_siguiente)
            else:
                btn = tk.Button(keypad, text=text, font=(FONT_PRIMARY, 14, "bold"), bg=COLOR_BG, fg=COLOR_TEXT_DARK, width=5, height=2, bd=1, relief="solid", cursor="hand2", command=lambda t=text: self.press_num(t))
            btn.grid(row=row, column=col, padx=4, pady=4)

        btn_volver = tk.Button(self, text="← REGRESAR AL MENÚ", font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_SECONDARY, fg="#FFFFFF", bd=0, relief="flat", padx=15, pady=8, cursor="hand2", command=lambda: controller.show_frame("PantallaPrincipal"))
        btn_volver.pack(pady=10)

    def on_show(self):
        self.display_var.set("")
        self.lbl_msg.config(text="")
        self.lbl_accion.config(text=f"OPERACIÓN: {self.controller.current_action} — INGRESE CÓDIGO DE COLABORADOR")

    def press_num(self, num):
        self.lbl_msg.config(text="")
        curr = self.display_var.get()
        if len(curr) < 8:
            self.display_var.set(curr + num)

    def press_borrar(self):
        self.lbl_msg.config(text="")
        self.display_var.set(self.display_var.get()[:-1])

    def press_siguiente(self):
        codigo = self.display_var.get().strip()
        if not codigo:
            return
        
        existe, depto = existe_empleado(codigo)
        if not existe:
            self.lbl_msg.config(text="⚠️ El código ingresado no está registrado en el sistema.", fg=COLOR_DANGER)
            return

        self.controller.current_emp_code = codigo
        self.controller.show_frame("PantallaBiometrica")


# ----------------------------------------------------------------------
# VISTA: RECONOCIMIENTO BIOMÉTRICO CON PRUEBA DE VIDA
# ----------------------------------------------------------------------
class PantallaBiometrica(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg=COLOR_BG)
        self.controller = controller
        self.cap = None
        self.is_running = False
        self.vector_registrado = None
        self.depto_empleado = ""
        self.verificado = False

        self.input_queue = mp.Queue()
        self.output_queue = mp.Queue()
        self.process = None

        self.last_face_rect = None
        self.last_color_box = (217, 130, 43)
        self.is_worker_busy = False

        HeaderNutresa(self).pack(fill="x")

        self.lbl_sub = tk.Label(self, text="Mire a la cámara y PARPADEE para verificar", font=(FONT_PRIMARY, 11, "bold"), fg=COLOR_MUTED, bg=COLOR_BG)
        self.lbl_sub.pack(pady=5)

        self.card = tk.Frame(self, bg=COLOR_CARD, bd=1, relief="solid")
        self.card.pack(pady=5, padx=80)

        self.lbl_video = tk.Label(self.card, bg="#000000", width=440, height=300)
        self.lbl_video.pack(pady=10, padx=10)

        self.lbl_status = tk.Label(self.card, text="INICIANDO CÁMARA...", font=(FONT_PRIMARY, 10, "bold"), fg=COLOR_MUTED, bg=COLOR_CARD)
        self.lbl_status.pack(pady=(0, 10))

        self.btn_cancelar = tk.Button(self, text="CANCELAR", font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_SECONDARY, fg="#FFFFFF", bd=0, relief="flat", padx=15, pady=8, cursor="hand2", command=lambda: controller.show_frame("PantallaPrincipal"))
        self.btn_cancelar.pack(pady=10)

    def on_show(self):
        self.verificado = False
        self.last_face_rect = None
        self.is_worker_busy = False

        emp_id = self.controller.current_emp_code
        ruta_vector = os.path.join(DIR_ROSTROS, f"{emp_id}.npy")

        if not os.path.exists(ruta_vector):
            self.lbl_sub.config(text=f"COLABORADOR ID: {emp_id}")
            self.lbl_status.config(text=f"❌ ERROR BIOMÉTRICO: El ID {emp_id} no tiene un rostro registrado.", fg=COLOR_DANGER)
            self.btn_cancelar.pack(pady=10)
            return

        _, self.depto_empleado = existe_empleado(emp_id)
        self.vector_registrado = np.load(ruta_vector)
        self.lbl_sub.config(text=f"COLABORADOR ID: {emp_id} — PARPADEE ANTE LA CÁMARA")
        self.lbl_status.config(text="👁 ESPERANDO PARPADEO DE VALIDACIÓN...", fg=COLOR_WARNING)
        self.btn_cancelar.pack(pady=10)

        self.detener_camara()

        self.input_queue = mp.Queue()
        self.output_queue = mp.Queue()
        self.process = mp.Process(
            target=worker_reconocimiento, 
            args=(self.input_queue, self.output_queue), 
            daemon=True
        )
        self.process.start()

        self.cap = cv2.VideoCapture(0)
        self.is_running = True
        self.actualizar_camara()

    def actualizar_camara(self):
        if not self.is_running or self.cap is None or not self.cap.isOpened():
            return

        ret, frame = self.cap.read()
        if ret:
            frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)

            if not self.is_worker_busy and not self.verificado:
                self.is_worker_busy = True
                try:
                    self.input_queue.put_nowait((frame_rgb, self.vector_registrado))
                except Exception:
                    self.is_worker_busy = False

            if not self.output_queue.empty():
                res = self.output_queue.get()
                self.is_worker_busy = False
                
                self.last_face_rect = res["rect"]

                if res["alerta_liveness"]:
                    self.last_color_box = (217, 130, 43)
                    self.lbl_status.config(text=res["alerta_liveness"], fg=COLOR_WARNING)
                elif res["verificado"]:
                    self.last_color_box = (0, 128, 55)
                    self.procesar_exito(res["similitud"])
                elif res["rect"] is not None:
                    self.last_color_box = (217, 56, 58)
                    self.lbl_status.config(
                        text=f"❌ ROSTRO NO COINCIDE ({res['similitud']:.1f}% Similitud)", 
                        fg=COLOR_DANGER
                    )

            if self.last_face_rect is not None:
                x1, y1, x2, y2 = self.last_face_rect
                cv2.rectangle(frame_rgb, (x1, y1), (x2, y2), self.last_color_box, 3)

            img = Image.fromarray(frame_rgb)
            img = img.resize((440, 300), Image.Resampling.LANCZOS)
            imgtk = ImageTk.PhotoImage(image=img)
            self.lbl_video.imgtk = imgtk
            self.lbl_video.configure(image=imgtk)

        if not self.verificado and self.is_running:
            self.after(10, self.actualizar_camara)

    def procesar_exito(self, similitud):
        self.verificado = True
        emp_id = self.controller.current_emp_code
        accion = self.controller.current_action

        self.lbl_status.config(
            text=f"✓ ¡PERSONA REAL VERIFICADA!\nCOLABORADOR ID: {emp_id}\nMARCACIÓN DE {accion} REGISTRADA", 
            fg=COLOR_PRIMARY
        )
        self.btn_cancelar.pack_forget()

        if accion in ["ENTRADA", "SALIDA", "TIEMPO EXTRA"]:
            registrar_marcacion_excel(emp_id, accion)
            self.after(1800, lambda: self.controller.show_frame("PantallaPrincipal"))
        elif accion == "MARCACIONES":
            self.after(1800, lambda: self.controller.show_frame("PantallaHistorial"))

    def detener_camara(self):
        self.is_running = False
        if self.cap is not None:
            self.cap.release()
            self.cap = None

        if self.process and self.process.is_alive():
            try:
                self.input_queue.put(None)
                self.process.terminate()
            except Exception:
                pass
            self.process = None


# ----------------------------------------------------------------------
# VISTA: CONSULTAR HISTORIAL DE MARCACIONES
# ----------------------------------------------------------------------
class PantallaHistorial(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg=COLOR_BG)
        self.controller = controller

        HeaderNutresa(self).pack(fill="x")

        self.lbl_titulo = tk.Label(self, text="", font=(FONT_PRIMARY, 11, "bold"), fg=COLOR_PRIMARY, bg=COLOR_BG)
        self.lbl_titulo.pack(pady=5)

        card_table = tk.Frame(self, bg=COLOR_CARD, bd=1, relief="solid")
        card_table.pack(pady=5, padx=30, fill="both", expand=True)

        columns = ("fecha", "entrada", "salida", "horas", "tarde")
        self.tree = ttk.Treeview(card_table, columns=columns, show="headings", height=10)
        self.tree.heading("fecha", text="FECHA")
        self.tree.heading("entrada", text="HORA ENTRADA")
        self.tree.heading("salida", text="HORA SALIDA")
        self.tree.heading("horas", text="HORAS TRAB.")
        self.tree.heading("tarde", text="MIN. TARDE")

        self.tree.column("fecha", anchor="center", width=120)
        self.tree.column("entrada", anchor="center", width=120)
        self.tree.column("salida", anchor="center", width=120)
        self.tree.column("horas", anchor="center", width=110)
        self.tree.column("tarde", anchor="center", width=110)

        scrollbar = ttk.Scrollbar(card_table, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscroll=scrollbar.set)

        self.tree.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        scrollbar.pack(side="right", fill="y", pady=10)

        btn_volver = tk.Button(self, text="← VOLVER AL MENÚ PRINCIPAL", font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_SECONDARY, fg="#FFFFFF", bd=0, relief="flat", padx=15, pady=8, cursor="hand2", command=lambda: controller.show_frame("PantallaPrincipal"))
        btn_volver.pack(pady=10)

    def on_show(self):
        emp_id = self.controller.current_emp_code
        _, depto = existe_empleado(emp_id)
        self.lbl_titulo.config(text=f"HISTORIAL DE MARCACIONES — COLABORADOR ID: {emp_id} ({depto})")

        for item in self.tree.get_children():
            self.tree.delete(item)

        registros = obtener_historial_empleado(emp_id)
        for reg in reversed(registros):
            self.tree.insert("", "end", values=(reg["fecha"], reg["entrada"], reg["salida"], reg["horas"], reg["tarde"]))


# ----------------------------------------------------------------------
# VISTA: ENROLAR NUEVO COLABORADOR (CON TECLADO TÁCTIL)
# ----------------------------------------------------------------------
class PantallaEnrolar(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg=COLOR_BG)
        self.controller = controller
        self.cap = None
        self.is_running = False
        self.ultimo_frame = None

        HeaderNutresa(self).pack(fill="x")

        lbl_titulo = tk.Label(self, text="ENROLAMIENTO FACIAL DE NUEVO COLABORADOR", font=(FONT_PRIMARY, 11, "bold"), fg=COLOR_PRIMARY, bg=COLOR_BG)
        lbl_titulo.pack(pady=2)

        main_frame = tk.Frame(self, bg=COLOR_BG)
        main_frame.pack(fill="both", expand=True, padx=20, pady=5)

        # Panel izquierdo (Formulario Táctil)
        panel_form = tk.Frame(main_frame, bg=COLOR_CARD, bd=1, relief="solid")
        panel_form.pack(side="left", fill="both", expand=True, padx=(5, 5), pady=5)

        tk.Label(panel_form, text="1. ID / CÉDULA:", font=(FONT_PRIMARY, 9, "bold"), fg=COLOR_MUTED, bg=COLOR_CARD).pack(anchor="w", padx=15, pady=(5, 0))
        
        self.var_id = tk.StringVar()
        self.ent_id = tk.Entry(panel_form, textvariable=self.var_id, font=(FONT_PRIMARY, 12, "bold"), justify="center", bd=1, relief="solid", state="readonly", bg=COLOR_BG)
        self.ent_id.pack(fill="x", padx=15, pady=(2, 5))

        keypad_frame = tk.Frame(panel_form, bg=COLOR_CARD)
        keypad_frame.pack(pady=2)

        buttons = [
            ('7', 0, 0), ('8', 0, 1), ('9', 0, 2),
            ('4', 1, 0), ('5', 1, 1), ('6', 1, 2),
            ('1', 2, 0), ('2', 2, 1), ('3', 2, 2),
            ('BORRAR', 3, 0), ('0', 3, 1)
        ]

        for text, row, col in buttons:
            if text == 'BORRAR':
                btn = tk.Button(keypad_frame, text=text, font=(FONT_PRIMARY, 8, "bold"), bg=COLOR_DANGER, fg="#FFFFFF", width=7, height=1, bd=0, relief="flat", cursor="hand2", command=self._borrar_id)
            else:
                btn = tk.Button(keypad_frame, text=text, font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_BG, fg=COLOR_TEXT_DARK, width=4, height=1, bd=1, relief="solid", cursor="hand2", command=lambda t=text: self._press_num(t))
            btn.grid(row=row, column=col, padx=2, pady=2)

        tk.Label(panel_form, text="2. DEPARTAMENTO:", font=(FONT_PRIMARY, 9, "bold"), fg=COLOR_MUTED, bg=COLOR_CARD).pack(anchor="w", padx=15, pady=(8, 0))
        self.var_depto = tk.StringVar(value="OPERACIONES")
        deptos = ["OPERACIONES", "LOGÍSTICA", "ADMINISTRACIÓN", "VENTAS", "MANTENIMIENTO"]
        opt_depto = ttk.OptionMenu(panel_form, self.var_depto, deptos[0], *deptos)
        opt_depto.pack(fill="x", padx=15, pady=2)

        self.lbl_status = tk.Label(panel_form, text="", font=(FONT_PRIMARY, 9, "bold"), bg=COLOR_CARD, wraplength=200)
        self.lbl_status.pack(pady=5)

        # Panel derecho (Cámara + Botón Guardar)
        panel_cam = tk.Frame(main_frame, bg=COLOR_CARD, bd=1, relief="solid")
        panel_cam.pack(side="right", fill="both", expand=True, padx=(5, 5), pady=5)

        self.lbl_video = tk.Label(panel_cam, bg="#000000", width=360, height=240)
        self.lbl_video.pack(pady=10, padx=10)

        btn_guardar = tk.Button(panel_cam, text="📸 CAPTURAR Y ENROLAR", font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_PRIMARY, fg="#FFFFFF", activebackground=COLOR_PRIMARY_HOVER, bd=0, relief="flat", cursor="hand2", command=self.guardar_enrolamiento)
        btn_guardar.pack(pady=5, ipadx=10, ipady=5)

        btn_volver = tk.Button(self, text="← REGRESAR AL MENÚ", font=(FONT_PRIMARY, 10, "bold"), bg=COLOR_SECONDARY, fg="#FFFFFF", bd=0, relief="flat", padx=15, pady=6, cursor="hand2", command=lambda: controller.show_frame("PantallaPrincipal"))
        btn_volver.pack(pady=(0, 10))

    def on_show(self):
        self.var_id.set("")
        self.lbl_status.config(text="")
        self.ultimo_frame = None
        self.detener_camara()
        self.cap = cv2.VideoCapture(0)
        self.is_running = True
        self.actualizar_camara()

    def _press_num(self, num):
        curr = self.var_id.get()
        if len(curr) < 8:
            self.var_id.set(curr + num)

    def _borrar_id(self):
        self.var_id.set(self.var_id.get()[:-1])

    def actualizar_camara(self):
        if not self.is_running or self.cap is None or not self.cap.isOpened():
            return

        ret, frame = self.cap.read()
        if ret:
            frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            self.ultimo_frame = frame_rgb.copy()

            locations = face_recognition.face_locations(cv2.resize(frame_rgb, (0, 0), fx=0.5, fy=0.5))
            if len(locations) > 0:
                top, right, bottom, left = locations[0]
                cv2.rectangle(frame_rgb, (left*2, top*2), (right*2, bottom*2), (0, 255, 0), 2)

            img = Image.fromarray(frame_rgb)
            img = img.resize((360, 240), Image.Resampling.LANCZOS)
            imgtk = ImageTk.PhotoImage(image=img)
            self.lbl_video.imgtk = imgtk
            self.lbl_video.configure(image=imgtk)

        if self.is_running:
            self.after(15, self.actualizar_camara)

    def guardar_enrolamiento(self):
        emp_id = self.var_id.get().strip()
        depto = self.var_depto.get().strip()

        if not emp_id:
            self.lbl_status.config(text="⚠️ Ingrese un ID válido.", fg=COLOR_DANGER)
            return

        existe, _ = existe_empleado(emp_id)
        if existe:
            self.lbl_status.config(text=f"⚠️ El ID {emp_id} ya existe.", fg=COLOR_DANGER)
            return

        if self.ultimo_frame is None:
            self.lbl_status.config(text="⚠️ Sin señal de cámara.", fg=COLOR_DANGER)
            return

        encodings = face_recognition.face_encodings(self.ultimo_frame)
        if len(encodings) == 0:
            self.lbl_status.config(text="❌ No se detectó rostro.", fg=COLOR_DANGER)
            return

        encoding = encodings[0]
        ruta_vector = os.path.join(DIR_ROSTROS, f"{emp_id}.npy")
        np.save(ruta_vector, encoding)

        guardar_nuevo_empleado(emp_id, depto, ruta_vector)
        self.lbl_status.config(text=f"✓ Colaborador {emp_id} enrolado con éxito.", fg=COLOR_PRIMARY)
        self.after(1500, lambda: self.controller.show_frame("PantallaPrincipal"))

    def detener_camara(self):
        self.is_running = False
        if self.cap is not None:
            self.cap.release()
            self.cap = None


# ----------------------------------------------------------------------
# PUNTO DE ENTRADA PRINCIPAL
# ----------------------------------------------------------------------
if __name__ == "__main__":
    mp.freeze_support()
    app = SistemaMarcacionApp()
    app.mainloop()